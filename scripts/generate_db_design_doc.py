"""Generate MAT Proposed Database Design (Common + MAT) Word + SQL DDL."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = Path(r"d:\MAT Project\MAT Doc")
DOCX_PATH = OUT_DIR / "MAT_Database_Design_Common_MAT_v2.docx"
SQL_PATH = OUT_DIR / "MAT_Database_Design_Common_MAT_v2.sql"


def set_cell_shading(cell, hex_color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tePr = cell._tc.get_or_add_tcPr()
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        set_cell_shading(cell, "D6EAF8")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


SQL_DDL = r"""
-- =============================================================================
-- Merchant Acquiring Tool (MAT) + Common (MAT/MEA)
-- Proposed Database Design v2 — Oracle DDL
-- Aligns with SRS maker-checker + admin flag model
-- =============================================================================
-- Conventions:
--   * All user references use EIN VARCHAR2(10) -> COMMON_USER.EIN
--   * IS_ADMIN on DCO users only (Y/N) — rights to Add user / Make admin
--   * Branch BU: AD login; DCO: must exist in COMMON_USER before login
--   * Maker <> Checker enforced in application (+ stored actor columns)
-- =============================================================================

-- ---------------------------------------------------------------------------
-- 0) SEQUENCES
-- ---------------------------------------------------------------------------
CREATE SEQUENCE SEQ_USER_APPROVAL START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_AUDIT_LOG START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_FILE_UPLOAD START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_VENDOR_TXN START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_REFUND_REQUEST START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_REQUEST_HISTORY START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_SCHEDULER_LOG START WITH 1 INCREMENT BY 1 NOCACHE;
CREATE SEQUENCE SEQ_EMAIL_LOG START WITH 1 INCREMENT BY 1 NOCACHE;

-- =============================================================================
-- COMMON SCHEMA OBJECTS (shared MAT + MEA)
-- =============================================================================

-- ---------------------------------------------------------------------------
-- 1) COMMON_ROLE  (was IDBI_ROLE)
-- ---------------------------------------------------------------------------
CREATE TABLE COMMON_ROLE (
    ROLE_ID       NUMBER(5)      NOT NULL,
    ROLE_CODE     VARCHAR2(50)   NOT NULL,  -- BU_USER, DCO_USER, INTERNAL_AUDITOR, EXTERNAL_AUDITOR, TEMP_USER
    ROLE_NAME     VARCHAR2(100)  NOT NULL,
    DESCRIPTION   VARCHAR2(200),
    IS_ACTIVE     CHAR(1)        DEFAULT 'Y' NOT NULL,
    CONSTRAINT PK_COMMON_ROLE PRIMARY KEY (ROLE_ID),
    CONSTRAINT UK_COMMON_ROLE_CODE UNIQUE (ROLE_CODE),
    CONSTRAINT CK_COMMON_ROLE_ACTIVE CHECK (IS_ACTIVE IN ('Y','N'))
);

COMMENT ON TABLE COMMON_ROLE IS 'Application roles shared by MAT and MEA';
COMMENT ON COLUMN COMMON_ROLE.ROLE_CODE IS 'BU_USER | DCO_USER | INTERNAL_AUDITOR | EXTERNAL_AUDITOR | TEMP_USER (and subtypes via USER_TYPE)';

-- Seed roles
INSERT INTO COMMON_ROLE (ROLE_ID, ROLE_CODE, ROLE_NAME, DESCRIPTION) VALUES (1, 'BU_USER', 'Branch User', 'Branch maker/checker via AD (RBG)');
INSERT INTO COMMON_ROLE (ROLE_ID, ROLE_CODE, ROLE_NAME, DESCRIPTION) VALUES (2, 'DCO_USER', 'DCO User', 'DCO review/approve; admin flag optional');
INSERT INTO COMMON_ROLE (ROLE_ID, ROLE_CODE, ROLE_NAME, DESCRIPTION) VALUES (3, 'INTERNAL_AUDITOR', 'Internal Auditor', 'Reports / view');
INSERT INTO COMMON_ROLE (ROLE_ID, ROLE_CODE, ROLE_NAME, DESCRIPTION) VALUES (4, 'EXTERNAL_AUDITOR', 'External Auditor', 'Reports / view');
INSERT INTO COMMON_ROLE (ROLE_ID, ROLE_CODE, ROLE_NAME, DESCRIPTION) VALUES (5, 'TEMP_USER', 'Temp User', 'BOA/HO/DBD/RO/ZO etc — time bound');

-- ---------------------------------------------------------------------------
-- 2) COMMON_USER  (was USER)
-- ---------------------------------------------------------------------------
CREATE TABLE COMMON_USER (
    EIN               VARCHAR2(10)   NOT NULL,          -- PK / AD username
    USERNAME          VARCHAR2(100)  NOT NULL,          -- Display name
    EMAIL             VARCHAR2(150),
    ROLE_ID           NUMBER(5)      NOT NULL,
    USER_TYPE         VARCHAR2(30),                     -- DCO_USER, INTERNAL_AUDITOR, EXTERNAL_AUDITOR, TEMP_BOA, TEMP_HO_DBD, ...
    SOL_ID            VARCHAR2(10),
    SOL_NAME          VARCHAR2(100),
    REGION_NAME       VARCHAR2(100),
    ZONE_NAME         VARCHAR2(100),
    POSITION_NAME     VARCHAR2(100),
    GRADE             VARCHAR2(50),
    ORGANIZATION      VARCHAR2(100),
    AGENCY            VARCHAR2(100),
    IS_ADMIN          CHAR(1)        DEFAULT 'N' NOT NULL, -- Y only meaningful for DCO_USER
    VALID_FROM        DATE,                               -- mandatory for non-DCO temp types
    VALID_TILL        DATE,                               -- INTERNAL/DCO can be far-future e.g. 31-DEC-2099
    IS_ACTIVATED      CHAR(1)        DEFAULT 'N' NOT NULL,
    ACTIVATED_DATE    DATE,
    ACTIVATED_BY      VARCHAR2(10),                       -- checker EIN
    IS_DEACTIVATED    CHAR(1)        DEFAULT 'N' NOT NULL,
    DEACTIVATED_DATE  DATE,
    DEACTIVATED_BY    VARCHAR2(10),
    IS_DELETED        CHAR(1)        DEFAULT 'N' NOT NULL, -- logical delete
    DELETED_DATE      DATE,
    DELETED_BY        VARCHAR2(10),
    CREATED_DATE      DATE           DEFAULT SYSDATE NOT NULL,
    CREATED_BY        VARCHAR2(10),                       -- maker EIN (or SYSTEM for seed)
    UPDATED_DATE      DATE,
    UPDATED_BY        VARCHAR2(10),
    CONSTRAINT PK_COMMON_USER PRIMARY KEY (EIN),
    CONSTRAINT FK_COMMON_USER_ROLE FOREIGN KEY (ROLE_ID) REFERENCES COMMON_ROLE (ROLE_ID),
    CONSTRAINT CK_COMMON_USER_ADMIN CHECK (IS_ADMIN IN ('Y','N')),
    CONSTRAINT CK_COMMON_USER_ACT CHECK (IS_ACTIVATED IN ('Y','N')),
    CONSTRAINT CK_COMMON_USER_DEA CHECK (IS_DEACTIVATED IN ('Y','N')),
    CONSTRAINT CK_COMMON_USER_DEL CHECK (IS_DELETED IN ('Y','N'))
);

CREATE INDEX IX_COMMON_USER_ROLE ON COMMON_USER (ROLE_ID);
CREATE INDEX IX_COMMON_USER_SOL ON COMMON_USER (SOL_ID);
CREATE INDEX IX_COMMON_USER_ADMIN ON COMMON_USER (IS_ADMIN);

COMMENT ON TABLE COMMON_USER IS 'Common users for MAT/MEA. Branch BU may be provisioned on first AD login; DCO must be pre-added.';
COMMENT ON COLUMN COMMON_USER.IS_ADMIN IS 'Y = can raise Add/Make-Admin user requests (DCO only). Does NOT block refund review/approve.';
COMMENT ON COLUMN COMMON_USER.CREATED_BY IS 'Maker EIN for user creation request (or SYSTEM for bootstrap admins)';
COMMENT ON COLUMN COMMON_USER.ACTIVATED_BY IS 'Checker EIN who approved activation';

-- ---------------------------------------------------------------------------
-- 3) COMMON_USER_APP_ACCESS  (MAT / MEA / both)
-- ---------------------------------------------------------------------------
CREATE TABLE COMMON_USER_APP_ACCESS (
    EIN         VARCHAR2(10)  NOT NULL,
    APP_CODE    VARCHAR2(10)  NOT NULL,   -- MAT | MEA
    IS_ACTIVE   CHAR(1)       DEFAULT 'Y' NOT NULL,
    GRANTED_DATE DATE         DEFAULT SYSDATE,
    GRANTED_BY  VARCHAR2(10),
    CONSTRAINT PK_USER_APP_ACCESS PRIMARY KEY (EIN, APP_CODE),
    CONSTRAINT FK_USER_APP_USER FOREIGN KEY (EIN) REFERENCES COMMON_USER (EIN),
    CONSTRAINT CK_USER_APP_CODE CHECK (APP_CODE IN ('MAT','MEA')),
    CONSTRAINT CK_USER_APP_ACTIVE CHECK (IS_ACTIVE IN ('Y','N'))
);

COMMENT ON TABLE COMMON_USER_APP_ACCESS IS 'Which applications a user may access. Toggle MAT/MEA uses this list.';

-- ---------------------------------------------------------------------------
-- 4) COMMON_USER_APPROVAL  (maker-checker for user management)
-- ---------------------------------------------------------------------------
CREATE TABLE COMMON_USER_APPROVAL (
    APPROVAL_ID       NUMBER(10)     NOT NULL,
    REQUEST_TYPE      VARCHAR2(20)   NOT NULL,  -- ADD, ACTIVATE, DEACTIVATE, DELETE, MAKE_ADMIN, REVOKE_ADMIN
    TARGET_EIN        VARCHAR2(10)   NOT NULL,  -- user being acted upon
    APP_CODE          VARCHAR2(10),             -- MAT/MEA/BOTH context if needed
    -- Snapshot for ADD (and display)
    TARGET_USERNAME   VARCHAR2(100),
    TARGET_EMAIL      VARCHAR2(150),
    TARGET_ROLE_ID    NUMBER(5),
    TARGET_USER_TYPE  VARCHAR2(30),
    TARGET_SOL_ID     VARCHAR2(10),
    TARGET_VALID_FROM DATE,
    TARGET_VALID_TILL DATE,
    REQUEST_ADMIN_YN  CHAR(1)        DEFAULT 'N', -- for ADD/MAKE_ADMIN: grant admin?
    STATUS            VARCHAR2(20)   DEFAULT 'PENDING' NOT NULL, -- PENDING, APPROVED, REJECTED
    MAKER_ID          VARCHAR2(10)   NOT NULL,
    MAKER_REMARK      VARCHAR2(500),
    MAKER_DATE        TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    CHECKER_ID        VARCHAR2(10),
    CHECKER_REMARK    VARCHAR2(500),
    CHECKER_DATE      TIMESTAMP,
    CONSTRAINT PK_USER_APPROVAL PRIMARY KEY (APPROVAL_ID),
    CONSTRAINT FK_UA_TARGET FOREIGN KEY (TARGET_EIN) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_UA_MAKER FOREIGN KEY (MAKER_ID) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_UA_CHECKER FOREIGN KEY (CHECKER_ID) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_UA_ROLE FOREIGN KEY (TARGET_ROLE_ID) REFERENCES COMMON_ROLE (ROLE_ID),
    CONSTRAINT CK_UA_TYPE CHECK (REQUEST_TYPE IN (
        'ADD','ACTIVATE','DEACTIVATE','DELETE','MAKE_ADMIN','REVOKE_ADMIN'
    )),
    CONSTRAINT CK_UA_STATUS CHECK (STATUS IN ('PENDING','APPROVED','REJECTED')),
    CONSTRAINT CK_UA_ADMIN_YN CHECK (REQUEST_ADMIN_YN IN ('Y','N')),
    CONSTRAINT CK_UA_MAKER_NE_CHECKER CHECK (CHECKER_ID IS NULL OR CHECKER_ID <> MAKER_ID)
);

CREATE INDEX IX_UA_STATUS ON COMMON_USER_APPROVAL (STATUS);
CREATE INDEX IX_UA_TARGET ON COMMON_USER_APPROVAL (TARGET_EIN);
CREATE INDEX IX_UA_MAKER ON COMMON_USER_APPROVAL (MAKER_ID);

COMMENT ON TABLE COMMON_USER_APPROVAL IS 'User management maker-checker. Maker cannot approve own request (CHECK constraint + app rule).';
COMMENT ON COLUMN COMMON_USER_APPROVAL.REQUEST_ADMIN_YN IS 'When ADD/MAKE_ADMIN approved, set COMMON_USER.IS_ADMIN=Y';

-- ---------------------------------------------------------------------------
-- 5) COMMON_AUDIT_LOG
-- ---------------------------------------------------------------------------
CREATE TABLE COMMON_AUDIT_LOG (
    AUDIT_ID      NUMBER(15)     NOT NULL,
    APP_CODE      VARCHAR2(10),                 -- MAT | MEA | COMMON
    USER_EIN      VARCHAR2(10),
    MODULE        VARCHAR2(50)   NOT NULL,
    ACTION        VARCHAR2(100)  NOT NULL,
    ENTITY_TYPE   VARCHAR2(50),                 -- REFUND_REQUEST, USER, FILE, ...
    ENTITY_ID     VARCHAR2(50),                 -- request id / ein / file id as string
    IP_ADDRESS    VARCHAR2(50),
    OLD_VALUE     CLOB,
    NEW_VALUE     CLOB,
    ACTION_TIME   TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    CONSTRAINT PK_AUDIT_LOG PRIMARY KEY (AUDIT_ID),
    CONSTRAINT FK_AUDIT_USER FOREIGN KEY (USER_EIN) REFERENCES COMMON_USER (EIN)
);

CREATE INDEX IX_AUDIT_USER ON COMMON_AUDIT_LOG (USER_EIN);
CREATE INDEX IX_AUDIT_ENTITY ON COMMON_AUDIT_LOG (ENTITY_TYPE, ENTITY_ID);
CREATE INDEX IX_AUDIT_TIME ON COMMON_AUDIT_LOG (ACTION_TIME);

-- =============================================================================
-- MAT SCHEMA OBJECTS
-- =============================================================================

-- ---------------------------------------------------------------------------
-- 6) MAT_FILE_UPLOAD
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_FILE_UPLOAD (
    FILE_ID         NUMBER(10)     NOT NULL,
    FILE_NAME       VARCHAR2(255)  NOT NULL,
    VENDOR          VARCHAR2(30)   NOT NULL,   -- WORLDLINE, HITACHI, SARVATRA
    FILE_TYPE       VARCHAR2(20)   NOT NULL,   -- POS, UPI, ARN
    SOURCE          VARCHAR2(20)   NOT NULL,   -- AUTO, MANUAL_SFTP, MANUAL_DESKTOP
    STATUS          VARCHAR2(20)   NOT NULL,   -- SUCCESS, FAILED, DELETED
    STORAGE_PATH    VARCHAR2(500),
    UPLOAD_TIME     TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    UPLOADED_BY     VARCHAR2(10),
    REMARK          VARCHAR2(500),
    CONSTRAINT PK_MAT_FILE PRIMARY KEY (FILE_ID),
    CONSTRAINT FK_MAT_FILE_USER FOREIGN KEY (UPLOADED_BY) REFERENCES COMMON_USER (EIN),
    CONSTRAINT CK_MAT_FILE_TYPE CHECK (FILE_TYPE IN ('POS','UPI','ARN')),
    CONSTRAINT CK_MAT_FILE_SOURCE CHECK (SOURCE IN ('AUTO','MANUAL_SFTP','MANUAL_DESKTOP'))
);

CREATE INDEX IX_MAT_FILE_VENDOR ON MAT_FILE_UPLOAD (VENDOR, FILE_TYPE, UPLOAD_TIME);

-- ---------------------------------------------------------------------------
-- 7) MAT_VENDOR_TRANSACTION  (one row from vendor file — source of truth for search)
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_VENDOR_TRANSACTION (
    TXN_ID            NUMBER(15)     NOT NULL,
    FILE_ID           NUMBER(10)     NOT NULL,
    VENDOR            VARCHAR2(30)   NOT NULL,
    TXN_CHANNEL       VARCHAR2(10)   NOT NULL,  -- UPI | POS
    CUSTOMER_ID       VARCHAR2(20),
    ACCOUNT_NO        VARCHAR2(20),
    RRN               VARCHAR2(30)   NOT NULL,
    MID               VARCHAR2(30)   NOT NULL,
    TID               VARCHAR2(30),
    MERCHANT_NAME     VARCHAR2(100),
    TXN_DATE          DATE           NOT NULL,
    TXN_AMOUNT        NUMBER(15,2)   NOT NULL,
    CARD_NUMBER_ENC   VARCHAR2(200),              -- encrypted
    CARD_NUMBER_MASK  VARCHAR2(30),
    AUTH_CODE_ENC     VARCHAR2(200),
    SCHEME            VARCHAR2(20),
    MERCHANT_VPA      VARCHAR2(100),
    GATEWAY_TXN_ID    VARCHAR2(50),
    ARN               VARCHAR2(50),
    IS_DELETED        CHAR(1)        DEFAULT 'N' NOT NULL,
    CREATED_DATE      TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    CONSTRAINT PK_MAT_VENDOR_TXN PRIMARY KEY (TXN_ID),
    CONSTRAINT FK_MAT_VENDOR_FILE FOREIGN KEY (FILE_ID) REFERENCES MAT_FILE_UPLOAD (FILE_ID),
    CONSTRAINT CK_MAT_VT_CHANNEL CHECK (TXN_CHANNEL IN ('UPI','POS')),
    CONSTRAINT CK_MAT_VT_DEL CHECK (IS_DELETED IN ('Y','N'))
);

CREATE INDEX IX_MAT_VT_LOOKUP ON MAT_VENDOR_TRANSACTION (RRN, TXN_DATE, TXN_AMOUNT);
CREATE INDEX IX_MAT_VT_MID ON MAT_VENDOR_TRANSACTION (MID);
CREATE INDEX IX_MAT_VT_FILE ON MAT_VENDOR_TRANSACTION (FILE_ID);

COMMENT ON TABLE MAT_VENDOR_TRANSACTION IS 'Parsed vendor SFTP/manual file rows used to validate Raise Request';

-- ---------------------------------------------------------------------------
-- 8) MAT_REFUND_REQUEST
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_REFUND_REQUEST (
    REQUEST_ID            NUMBER(30)     NOT NULL,
    REQUEST_TYPE          VARCHAR2(10)   NOT NULL,  -- UPI | POS
    TXN_ID                NUMBER(15),                 -- FK vendor txn (nullable if exceptional)
    CUSTOMER_ID           VARCHAR2(20)   NOT NULL,
    ACCOUNT_NO            VARCHAR2(20)   NOT NULL,
    BRANCH_SOL            VARCHAR2(10)   NOT NULL,
    RRN                   VARCHAR2(30)   NOT NULL,
    MID                   VARCHAR2(30)   NOT NULL,
    TID                   VARCHAR2(30),
    TXN_DATE              DATE           NOT NULL,
    TXN_AMOUNT            NUMBER(15,2)   NOT NULL,
    REFUND_AMOUNT         NUMBER(15,2)   NOT NULL,
    CARD_NUMBER_MASK      VARCHAR2(30),
    AUTH_CODE_MASK        VARCHAR2(30),
    SCHEME                VARCHAR2(20),
    MERCHANT_VPA          VARCHAR2(100),
    GATEWAY_TXN_ID        VARCHAR2(50),
    VENDOR                VARCHAR2(30),
    REFERENCE_DOC_PATH    VARCHAR2(500),
    -- Workflow
    STATUS                VARCHAR2(40)   NOT NULL,
    -- Examples: DRAFT, BRANCH_PENDING, BRANCH_APPROVED, BRANCH_REJECTED, REFERRED_BACK,
    --           DCO_REVIEW_PENDING, DCO_REVIEWED, DCO_APPROVED, DCO_REJECTED,
    --           VENDOR_REJECTED, ARN_PENDING, TXN_SUCCESS, TXN_FAILED, DELETED, ARCHIVED
    STAGE                 VARCHAR2(30),               -- BRANCH | DCO_REVIEW | DCO_APPROVE | VENDOR | FINACLE | CLOSED
    RETRY_COUNT           NUMBER(2)      DEFAULT 0,
    VENDOR_REJECT_REASON  VARCHAR2(500),
    -- Branch maker / checker
    CREATED_BY            VARCHAR2(10)   NOT NULL,   -- Branch maker EIN
    CREATED_DATE          TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    BRANCH_APPROVED_BY    VARCHAR2(10),               -- must <> CREATED_BY
    BRANCH_APPROVED_DATE  TIMESTAMP,
    BRANCH_REMARK         VARCHAR2(500),
    -- DCO review / final approve
    DCO_REVIEWED_BY       VARCHAR2(10),               -- reviewer
    DCO_REVIEWED_DATE     TIMESTAMP,
    DCO_REVIEW_REMARK     VARCHAR2(500),
    DCO_APPROVED_BY       VARCHAR2(10),               -- final; must <> DCO_REVIEWED_BY
    DCO_APPROVED_DATE     TIMESTAMP,
    DCO_APPROVE_REMARK    VARCHAR2(500),
    UPDATED_BY            VARCHAR2(10),
    UPDATED_DATE          TIMESTAMP,
    IS_DELETED            CHAR(1)        DEFAULT 'N' NOT NULL,
    CONSTRAINT PK_MAT_REFUND PRIMARY KEY (REQUEST_ID),
    CONSTRAINT FK_MAT_REFUND_TXN FOREIGN KEY (TXN_ID) REFERENCES MAT_VENDOR_TRANSACTION (TXN_ID),
    CONSTRAINT FK_MAT_REFUND_MAKER FOREIGN KEY (CREATED_BY) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_MAT_REFUND_BAPP FOREIGN KEY (BRANCH_APPROVED_BY) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_MAT_REFUND_DREV FOREIGN KEY (DCO_REVIEWED_BY) REFERENCES COMMON_USER (EIN),
    CONSTRAINT FK_MAT_REFUND_DAPP FOREIGN KEY (DCO_APPROVED_BY) REFERENCES COMMON_USER (EIN),
    CONSTRAINT CK_MAT_REFUND_TYPE CHECK (REQUEST_TYPE IN ('UPI','POS')),
    CONSTRAINT CK_MAT_REFUND_DEL CHECK (IS_DELETED IN ('Y','N')),
    CONSTRAINT CK_MAT_REFUND_BRANCH_NE CHECK (
        BRANCH_APPROVED_BY IS NULL OR BRANCH_APPROVED_BY <> CREATED_BY
    ),
    CONSTRAINT CK_MAT_REFUND_DCO_NE CHECK (
        DCO_APPROVED_BY IS NULL OR DCO_REVIEWED_BY IS NULL OR DCO_APPROVED_BY <> DCO_REVIEWED_BY
    )
);

CREATE INDEX IX_MAT_REFUND_STATUS ON MAT_REFUND_REQUEST (STATUS, STAGE);
CREATE INDEX IX_MAT_REFUND_MAKER ON MAT_REFUND_REQUEST (CREATED_BY);
CREATE INDEX IX_MAT_REFUND_DUP ON MAT_REFUND_REQUEST (RRN, TXN_DATE, REFUND_AMOUNT);
CREATE INDEX IX_MAT_REFUND_SOL ON MAT_REFUND_REQUEST (BRANCH_SOL);

COMMENT ON TABLE MAT_REFUND_REQUEST IS 'MAT refund/reversal requests with Branch + DCO maker-checker actor columns';
COMMENT ON COLUMN MAT_REFUND_REQUEST.CREATED_BY IS 'Branch maker — excluded from own Approval Queue';
COMMENT ON COLUMN MAT_REFUND_REQUEST.DCO_REVIEWED_BY IS 'DCO reviewer — cannot be final approver for same request';
COMMENT ON COLUMN MAT_REFUND_REQUEST.DCO_APPROVED_BY IS 'DCO final approver — any DCO (admin or not)';

-- ---------------------------------------------------------------------------
-- 9) MAT_REQUEST_HISTORY
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_REQUEST_HISTORY (
    HISTORY_ID    NUMBER(15)     NOT NULL,
    REQUEST_ID    NUMBER(30)     NOT NULL,
    ACTION        VARCHAR2(50)   NOT NULL,  -- CREATE, BRANCH_APPROVE, BRANCH_REJECT, REFER_BACK, DCO_REVIEW, DCO_APPROVE, ...
    OLD_STATUS    VARCHAR2(40),
    NEW_STATUS    VARCHAR2(40),
    REMARK        VARCHAR2(500),
    ACTION_BY     VARCHAR2(10)   NOT NULL,
    ACTION_DATE   TIMESTAMP      DEFAULT SYSTIMESTAMP NOT NULL,
    CONSTRAINT PK_MAT_REQ_HIST PRIMARY KEY (HISTORY_ID),
    CONSTRAINT FK_MAT_HIST_REQ FOREIGN KEY (REQUEST_ID) REFERENCES MAT_REFUND_REQUEST (REQUEST_ID),
    CONSTRAINT FK_MAT_HIST_USER FOREIGN KEY (ACTION_BY) REFERENCES COMMON_USER (EIN)
);

CREATE INDEX IX_MAT_HIST_REQ ON MAT_REQUEST_HISTORY (REQUEST_ID, ACTION_DATE);

-- ---------------------------------------------------------------------------
-- 10) MAT_SCHEDULER_LOG
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_SCHEDULER_LOG (
    JOB_ID        NUMBER(10)     NOT NULL,
    JOB_NAME      VARCHAR2(100)  NOT NULL,
    START_TIME    TIMESTAMP      NOT NULL,
    END_TIME      TIMESTAMP,
    STATUS        VARCHAR2(20)   NOT NULL,  -- SUCCESS, FAILED, RUNNING
    MESSAGE       VARCHAR2(1000),
    CONSTRAINT PK_MAT_SCHED PRIMARY KEY (JOB_ID)
);

-- ---------------------------------------------------------------------------
-- 11) MAT_EMAIL_LOG (optional)
-- ---------------------------------------------------------------------------
CREATE TABLE MAT_EMAIL_LOG (
    EMAIL_ID      NUMBER(10)     NOT NULL,
    REQUEST_ID    NUMBER(30),
    TO_EMAIL      VARCHAR2(200)  NOT NULL,
    CC_EMAIL      VARCHAR2(500),
    SUBJECT       VARCHAR2(300)  NOT NULL,
    STATUS        VARCHAR2(20)   NOT NULL,  -- SENT, FAILED
    SENT_TIME     TIMESTAMP      DEFAULT SYSTIMESTAMP,
    ERROR_MSG     VARCHAR2(1000),
    CONSTRAINT PK_MAT_EMAIL PRIMARY KEY (EMAIL_ID),
    CONSTRAINT FK_MAT_EMAIL_REQ FOREIGN KEY (REQUEST_ID) REFERENCES MAT_REFUND_REQUEST (REQUEST_ID)
);

-- =============================================================================
-- BOOTSTRAP: 2 DCO admin users (replace EINs with ops-provided values)
-- =============================================================================
-- INSERT INTO COMMON_USER (EIN, USERNAME, EMAIL, ROLE_ID, USER_TYPE, SOL_ID, IS_ADMIN,
--   VALID_TILL, IS_ACTIVATED, ACTIVATED_DATE, ACTIVATED_BY, CREATED_BY, CREATED_DATE)
-- VALUES ('100001', 'DCO Admin One', 'dco1@idbi.co.in', 2, 'DCO_USER', '999', 'Y',
--   DATE '2099-12-31', 'Y', SYSDATE, 'SYSTEM', 'SYSTEM', SYSDATE);
--
-- INSERT INTO COMMON_USER (EIN, USERNAME, EMAIL, ROLE_ID, USER_TYPE, SOL_ID, IS_ADMIN,
--   VALID_TILL, IS_ACTIVATED, ACTIVATED_DATE, ACTIVATED_BY, CREATED_BY, CREATED_DATE)
-- VALUES ('100002', 'DCO Admin Two', 'dco2@idbi.co.in', 2, 'DCO_USER', '999', 'Y',
--   DATE '2099-12-31', 'Y', SYSDATE, 'SYSTEM', 'SYSTEM', SYSDATE);
--
-- INSERT INTO COMMON_USER_APP_ACCESS (EIN, APP_CODE) VALUES ('100001','MAT');
-- INSERT INTO COMMON_USER_APP_ACCESS (EIN, APP_CODE) VALUES ('100001','MEA');
-- INSERT INTO COMMON_USER_APP_ACCESS (EIN, APP_CODE) VALUES ('100002','MAT');
-- INSERT INTO COMMON_USER_APP_ACCESS (EIN, APP_CODE) VALUES ('100002','MEA');

COMMIT;
"""


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_heading("Merchant Acquiring Tool (MAT)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Proposed Database Design v2 — Common + MAT (Oracle)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        "Version: 2.0  |  Date: 06-Aug-2026  |  Based on SRS v3 + maker maker-checker / admin model"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "This document updates the earlier MAT_Proposed_Database_Design with: "
        "(a) clear Common vs MAT table split for future MEA sharing; "
        "(b) IS_ADMIN flag for DCO user-management rights; "
        "(c) consistent EIN-based user foreign keys; "
        "(d) Branch + DCO maker-checker actor columns on refund requests; "
        "(e) enriched USER_APPROVAL for Add/Activate/Deactivate/Delete/Make-Admin."
    )

    doc.add_heading("2. Design principles", level=1)
    for item in [
        "COMMON_* tables are shared by MAT and MEA (users, roles, user approvals, audit).",
        "MAT_* tables are refund / vendor-file / scheduler specific.",
        "All user FKs reference COMMON_USER.EIN (VARCHAR2(10)) — no NUMBER user ids.",
        "DCO is one role (DCO_USER). Admin is IS_ADMIN='Y' on that user — not a separate role.",
        "Admin DCO can still raise/review/approve refunds; admin only adds User Management rights.",
        "Maker ≠ Checker enforced in DB checks where possible and always in application logic.",
        "Branch BU: AD login (RBG). DCO: must exist in COMMON_USER (SOL typically 999/9990/935).",
        "Bootstrap: seed two DCO users with IS_ADMIN='Y' from backend (ops-provided EINs).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. Logical model", level=1)
    doc.add_paragraph("COMMON")
    for t in [
        "COMMON_ROLE",
        "COMMON_USER",
        "COMMON_USER_APP_ACCESS",
        "COMMON_USER_APPROVAL",
        "COMMON_AUDIT_LOG",
    ]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph("MAT")
    for t in [
        "MAT_FILE_UPLOAD",
        "MAT_VENDOR_TRANSACTION",
        "MAT_REFUND_REQUEST",
        "MAT_REQUEST_HISTORY",
        "MAT_SCHEDULER_LOG",
        "MAT_EMAIL_LOG",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("4. Changes vs previous design", level=1)
    add_table(
        doc,
        ["Previous", "v2 Change", "Reason"],
        [
            ["IDBI_ROLE / USER", "COMMON_ROLE / COMMON_USER", "Shared MAT+MEA naming"],
            ["SYSTEM-ACCESS MAT/MEA", "COMMON_USER_APP_ACCESS", "Support both apps"],
            ["No IS_ADMIN", "COMMON_USER.IS_ADMIN", "Add user / make admin rights"],
            ["USER FKs as NUMBER", "All FKs = EIN VARCHAR2(10)", "PK/FK consistency"],
            ["USER_APPROVAL thin", "REQUEST_TYPE + remarks + snapshot", "Full user maker-checker"],
            ["REFUND_REQUEST no actors", "CREATED_BY / BRANCH_* / DCO_*", "SoD maker-checker"],
            ["VENDOR_TXN_INPUT_FILE_DATA mixed", "MAT_VENDOR_TRANSACTION clean", "Separate file row vs request"],
            ["DELETED _BY typo", "DELETED_BY", "Correct column name"],
            ["VALID_TILL only", "+ VALID_FROM", "Temp user time frame (SRS)"],
        ],
    )

    # Table specs
    doc.add_heading("5. Common tables", level=1)

    doc.add_heading("5.1 COMMON_ROLE", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["ROLE_ID", "NUMBER(5)", "PK", "Role id"],
            ["ROLE_CODE", "VARCHAR2(50)", "UNIQUE", "BU_USER, DCO_USER, INTERNAL_AUDITOR, EXTERNAL_AUDITOR, TEMP_USER"],
            ["ROLE_NAME", "VARCHAR2(100)", "", "Display name"],
            ["DESCRIPTION", "VARCHAR2(200)", "", "Description"],
            ["IS_ACTIVE", "CHAR(1)", "DEFAULT Y", "Y/N"],
        ],
    )

    doc.add_heading("5.2 COMMON_USER", level=2)
    doc.add_paragraph(
        "Note: IS_ADMIN='Y' is only used for User Management (Add / Make Admin). "
        "Admin DCO users can still review and approve refund requests like any DCO."
    )
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["EIN", "VARCHAR2(10)", "PK", "User id (AD)"],
            ["USERNAME", "VARCHAR2(100)", "", "Display name"],
            ["EMAIL", "VARCHAR2(150)", "", "Email"],
            ["ROLE_ID", "NUMBER(5)", "FK", "→ COMMON_ROLE"],
            ["USER_TYPE", "VARCHAR2(30)", "", "DCO_USER / auditor / TEMP_BOA / TEMP_HO_DBD …"],
            ["SOL_ID", "VARCHAR2(10)", "", "SOL (DCO: 999/9990/935)"],
            ["SOL_NAME / REGION_NAME / ZONE_NAME", "VARCHAR2(100)", "", "Org attributes"],
            ["POSITION_NAME / GRADE", "VARCHAR2", "", "HRMS"],
            ["IS_ADMIN", "CHAR(1)", "DEFAULT N", "Y = can add users / grant admin"],
            ["VALID_FROM / VALID_TILL", "DATE", "", "Mandatory for temp types"],
            ["IS_ACTIVATED / ACTIVATED_BY / DATE", "", "", "Checker activation"],
            ["IS_DEACTIVATED / DEACTIVATED_BY / DATE", "", "", "Checker deactivation"],
            ["IS_DELETED / DELETED_BY / DATE", "", "", "Logical delete"],
            ["CREATED_BY / CREATED_DATE", "", "", "Maker (or SYSTEM for seed)"],
        ],
    )

    doc.add_heading("5.3 COMMON_USER_APP_ACCESS", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["EIN", "VARCHAR2(10)", "PK/FK", "User"],
            ["APP_CODE", "VARCHAR2(10)", "PK", "MAT | MEA"],
            ["IS_ACTIVE", "CHAR(1)", "DEFAULT Y", "Access flag"],
            ["GRANTED_BY / GRANTED_DATE", "", "", "Audit"],
        ],
    )

    doc.add_heading("5.4 COMMON_USER_APPROVAL", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["APPROVAL_ID", "NUMBER(10)", "PK", "Sequence"],
            ["REQUEST_TYPE", "VARCHAR2(20)", "", "ADD | ACTIVATE | DEACTIVATE | DELETE | MAKE_ADMIN | REVOKE_ADMIN"],
            ["TARGET_EIN", "VARCHAR2(10)", "FK", "User acted upon"],
            ["TARGET_* snapshot", "various", "", "Name, email, role, type, SOL, validity"],
            ["REQUEST_ADMIN_YN", "CHAR(1)", "", "Grant admin on approve"],
            ["STATUS", "VARCHAR2(20)", "", "PENDING | APPROVED | REJECTED"],
            ["MAKER_ID / MAKER_REMARK / MAKER_DATE", "", "FK", "Initiator (must be admin for ADD/MAKE_ADMIN)"],
            ["CHECKER_ID / CHECKER_REMARK / CHECKER_DATE", "", "FK", "Must <> MAKER_ID"],
        ],
    )

    doc.add_heading("5.5 COMMON_AUDIT_LOG", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["AUDIT_ID", "NUMBER(15)", "PK", "Sequence"],
            ["APP_CODE", "VARCHAR2(10)", "", "MAT | MEA | COMMON"],
            ["USER_EIN", "VARCHAR2(10)", "FK", "Actor"],
            ["MODULE / ACTION", "VARCHAR2", "", "Module action"],
            ["ENTITY_TYPE / ENTITY_ID", "VARCHAR2", "", "Generic entity ref"],
            ["IP_ADDRESS", "VARCHAR2(50)", "", "Client IP"],
            ["OLD_VALUE / NEW_VALUE", "CLOB", "", "Change payload"],
            ["ACTION_TIME", "TIMESTAMP", "", "When"],
        ],
    )

    doc.add_heading("6. MAT tables", level=1)

    doc.add_heading("6.1 MAT_FILE_UPLOAD", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["FILE_ID", "NUMBER(10)", "PK", "Sequence"],
            ["FILE_NAME", "VARCHAR2(255)", "", "Name"],
            ["VENDOR", "VARCHAR2(30)", "", "WORLDLINE / HITACHI / SARVATRA"],
            ["FILE_TYPE", "VARCHAR2(20)", "", "POS | UPI | ARN"],
            ["SOURCE", "VARCHAR2(20)", "", "AUTO | MANUAL_SFTP | MANUAL_DESKTOP"],
            ["STATUS", "VARCHAR2(20)", "", "SUCCESS | FAILED | DELETED"],
            ["UPLOADED_BY", "VARCHAR2(10)", "FK", "EIN"],
            ["UPLOAD_TIME", "TIMESTAMP", "", "When"],
        ],
    )

    doc.add_heading("6.2 MAT_VENDOR_TRANSACTION", level=2)
    doc.add_paragraph(
        "Replaces mixed VENDOR_TXN_INPUT_FILE_DATA. Holds parsed vendor file rows only. "
        "Refund requests reference TXN_ID."
    )
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["TXN_ID", "NUMBER(15)", "PK", "Sequence"],
            ["FILE_ID", "NUMBER(10)", "FK", "→ MAT_FILE_UPLOAD"],
            ["VENDOR / TXN_CHANNEL", "VARCHAR2", "", "UPI | POS"],
            ["RRN / MID / TID / TXN_DATE / TXN_AMOUNT", "", "", "Lookup keys"],
            ["CARD_NUMBER_ENC / MASK", "VARCHAR2", "", "Encrypted + display mask"],
            ["AUTH_CODE_ENC / SCHEME / VPA / GATEWAY_TXN_ID / ARN", "", "", "Channel fields"],
        ],
    )

    doc.add_heading("6.3 MAT_REFUND_REQUEST", level=2)
    doc.add_paragraph(
        "Branch: CREATED_BY (maker) ≠ BRANCH_APPROVED_BY (checker). "
        "DCO: DCO_REVIEWED_BY ≠ DCO_APPROVED_BY. Admin flag not required for these actions."
    )
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["REQUEST_ID", "NUMBER(30)", "PK", "Request id"],
            ["REQUEST_TYPE", "VARCHAR2(10)", "", "UPI | POS"],
            ["TXN_ID", "NUMBER(15)", "FK", "Vendor txn link"],
            ["CUSTOMER_ID / ACCOUNT_NO / BRANCH_SOL", "", "", "CBS / branch"],
            ["RRN / MID / TID / TXN_DATE / TXN_AMOUNT / REFUND_AMOUNT", "", "", "Request amounts"],
            ["STATUS / STAGE", "VARCHAR2", "", "Lifecycle"],
            ["CREATED_BY / CREATED_DATE", "VARCHAR2(10)", "FK", "Branch maker"],
            ["BRANCH_APPROVED_BY / DATE / BRANCH_REMARK", "", "FK", "Branch checker ≠ maker"],
            ["DCO_REVIEWED_BY / DATE / REMARK", "", "FK", "DCO reviewer"],
            ["DCO_APPROVED_BY / DATE / REMARK", "", "FK", "DCO final ≠ reviewer"],
            ["VENDOR_REJECT_REASON / RETRY_COUNT / REFERENCE_DOC_PATH", "", "", "Ops fields"],
            ["IS_DELETED", "CHAR(1)", "DEFAULT N", "Logical delete"],
        ],
    )

    doc.add_heading("6.4 MAT_REQUEST_HISTORY", level=2)
    add_table(
        doc,
        ["Column", "Data Type", "Key", "Description"],
        [
            ["HISTORY_ID", "NUMBER(15)", "PK", "Sequence"],
            ["REQUEST_ID", "NUMBER(30)", "FK", "Refund request"],
            ["ACTION", "VARCHAR2(50)", "", "CREATE, BRANCH_APPROVE, DCO_REVIEW, …"],
            ["OLD_STATUS / NEW_STATUS", "VARCHAR2(40)", "", "Transition"],
            ["REMARK", "VARCHAR2(500)", "", "Remark"],
            ["ACTION_BY / ACTION_DATE", "VARCHAR2(10) / TIMESTAMP", "FK", "Actor EIN"],
        ],
    )

    doc.add_heading("6.5 MAT_SCHEDULER_LOG / MAT_EMAIL_LOG", level=2)
    doc.add_paragraph(
        "Unchanged in intent from prior design; user FKs (if any) use EIN. "
        "See accompanying .sql file for full column lists."
    )

    doc.add_heading("7. Status values (recommended)", level=1)
    doc.add_paragraph("Refund STATUS examples:", style="List Bullet")
    doc.add_paragraph(
        "DRAFT, BRANCH_PENDING, BRANCH_APPROVED, BRANCH_REJECTED, REFERRED_BACK, "
        "DCO_REVIEW_PENDING, DCO_REVIEWED, DCO_APPROVED, DCO_REJECTED, "
        "VENDOR_REJECTED, ARN_PENDING, TXN_SUCCESS, TXN_FAILED, DELETED, ARCHIVED"
    )
    doc.add_paragraph("User approval STATUS: PENDING | APPROVED | REJECTED")

    doc.add_heading("8. Application rules (enforce in API)", level=1)
    for item in [
        "Branch Approval Queue: exclude rows where CREATED_BY = current EIN.",
        "DCO final approve: reject if current EIN = DCO_REVIEWED_BY.",
        "User Management ADD/MAKE_ADMIN: maker must have IS_ADMIN='Y'.",
        "User Management checker: CHECKER_ID <> MAKER_ID; prefer another admin per SRS v3.",
        "DCO login: user must exist, IS_ACTIVATED='Y', IS_DELETED='N', SOL in allowed set, app access MAT.",
        "Branch login: AD success + RBG; provision/update COMMON_USER as BU_USER as per bank policy.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("9. Bootstrap", level=1)
    doc.add_paragraph(
        "Insert two DCO users with IS_ADMIN='Y', IS_ACTIVATED='Y', CREATED_BY='SYSTEM', "
        "and grant COMMON_USER_APP_ACCESS for MAT (and MEA if required). "
        "EINs to be provided by Operations. Sample INSERT statements are in the .sql file."
    )

    doc.add_heading("10. DDL script", level=1)
    doc.add_paragraph(
        f"Full Oracle DDL is provided alongside this document:\n{SQL_PATH.name}"
    )
    doc.add_paragraph(
        "Run sequences → common tables → MAT tables → seed roles → bootstrap admins."
    )

    doc.add_heading("11. Out of scope (MEA)", level=1)
    doc.add_paragraph(
        "MEA enrolment tables are not defined here. They will reference COMMON_USER / "
        "COMMON_ROLE / COMMON_USER_APP_ACCESS / COMMON_AUDIT_LOG when MEA build starts."
    )

    doc.add_paragraph("")
    end = doc.add_paragraph("— End of Document —")
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SQL_PATH.write_text(SQL_DDL.lstrip("\n"), encoding="utf-8")
    build_docx()
    print(f"Wrote {DOCX_PATH}")
    print(f"Wrote {SQL_PATH}")


if __name__ == "__main__":
    main()

"""Recreate MAT_Proposed_Database_Design.docx in the same simple format with updated tables."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path(r"d:\MAT Project\MAT Doc\MAT_Proposed_Database_Design_Updated.docx")


def shade_header(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_col_table(doc, rows):
    """rows: list of (column, datatype, key, description)"""
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    headers = ["Column", "Data Type", "Key", "Description"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
                r.font.name = "Calibri"
        shade_header(cell)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.name = "Calibri"
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    doc.add_heading("Merchant Acquiring Tool (MAT) - Proposed Database Design", level=1)

    # ----- IDBI ROLE -----
    doc.add_heading("IDBI ROLE", level=2)
    add_col_table(
        doc,
        [
            ("ROLE_ID", "NUMBER(5)", "PK", "Role ID"),
            (
                "ROLE_NAME",
                "VARCHAR2(50)",
                "UNIQUE",
                "DCO_USER, BU_USER, INTERNAL_AUDITOR, EXTERNAL_AUDITOR, HO/DBD/TEMP",
            ),
            ("DESCRIPTION", "VARCHAR2(200)", "", "Description"),
        ],
    )

    # ----- USER -----
    doc.add_heading("USER", level=2)
    add_col_table(
        doc,
        [
            ("EIN", "VARCHAR2(10)", "PK", "User ID (AD)"),
            ("USERNAME", "VARCHAR2(100)", "", "Name"),
            ("EMAIL", "VARCHAR2(150)", "", "Email"),
            ("ROLE_ID", "NUMBER(5)", "FK", "Role"),
            ("USER_TYPE", "VARCHAR2(30)", "", "DCO_USER / INTERNAL_AUDITOR / EXTERNAL_AUDITOR / TEMP_BOA / TEMP_HO_DBD etc."),
            ("SOL_ID", "VARCHAR2(10)", "", "SOL ID (DCO: 999 / 9990 / 935)"),
            ("SOL_NAME", "VARCHAR2(100)", "", "SOL Name"),
            ("REGION_NAME", "VARCHAR2(100)", "", "Region"),
            ("ZONE_NAME", "VARCHAR2(100)", "", "Zone"),
            ("POSITION", "VARCHAR2(100)", "", "Position"),
            ("GRADE", "VARCHAR2(50)", "", "Grade"),
            ("ORGANIZATION", "VARCHAR2(100)", "", "Organization"),
            ("AGENCY", "VARCHAR2(100)", "", "IDBI BANK LTD / ______________"),
            ("IS_ADMIN", "CHAR(1)", "DEFAULT N", "Y = rights to Add user / Make admin (DCO only). Admin can still raise/approve refunds."),
            ("VALID_FROM", "DATE", "", "From date (mandatory for temp users)"),
            ("VALID_TILL", "DATE", "", "Last valid date (for internal/DCO use 31-DEC-2099)"),
            ("CREATED_DATE", "DATE", "", ""),
            ("CREATED_BY", "VARCHAR2(10)", "FK", "Maker User EIN (or SYSTEM for seed admins)"),
            ("ACTIVATED_DATE", "DATE", "", ""),
            ("ACTIVATED_BY", "VARCHAR2(10)", "FK", "Checker User EIN"),
            ("IS_ACTIVATED", "CHAR(1)", "DEFAULT N", ""),
            ("DELETED_DATE", "DATE", "", ""),
            ("DELETED_BY", "VARCHAR2(10)", "FK", "Checker User EIN"),
            ("IS_DELETED", "CHAR(1)", "DEFAULT N", "Logical delete"),
            ("DEACTIVATED_DATE", "DATE", "", ""),
            ("DEACTIVATED_BY", "VARCHAR2(10)", "FK", "Checker User EIN"),
            ("IS_DEACTIVATED", "CHAR(1)", "DEFAULT N", ""),
        ],
    )

    # ----- USER_APP_ACCESS (new — replaces SYSTEM-ACCESS column) -----
    doc.add_heading("USER_APP_ACCESS", level=2)
    add_col_table(
        doc,
        [
            ("EIN", "VARCHAR2(10)", "PK, FK", "User EIN"),
            ("APP_CODE", "VARCHAR2(10)", "PK", "MAT / MEA"),
            ("IS_ACTIVE", "CHAR(1)", "DEFAULT Y", "Access flag"),
            ("GRANTED_DATE", "DATE", "", ""),
            ("GRANTED_BY", "VARCHAR2(10)", "FK", "Granted by EIN"),
        ],
    )

    # ----- VENDOR_TRANSACTION (clean file rows; replaces mixed INPUT_FILE_DATA) -----
    doc.add_heading("VENDOR_TRANSACTION", level=2)
    add_col_table(
        doc,
        [
            ("TXN_ID", "NUMBER(15)", "PK", ""),
            ("FILE_ID", "NUMBER(10)", "FK", "FILE_UPLOAD"),
            ("VENDOR", "VARCHAR2(30)", "", "WORLDLINE / HITACHI / SARVATRA"),
            ("TXN_CHANNEL", "VARCHAR2(10)", "", "UPI / POS"),
            ("CUSTOMER_ID", "VARCHAR2(20)", "", ""),
            ("ACCOUNT_NO", "VARCHAR2(20)", "", ""),
            ("RRN", "VARCHAR2(30)", "", ""),
            ("MID", "VARCHAR2(30)", "", ""),
            ("TID", "VARCHAR2(30)", "", ""),
            ("MERCHANT_NAME", "VARCHAR2(100)", "", ""),
            ("TXN_DATE", "DATE", "", ""),
            ("TXN_AMOUNT", "NUMBER(15,2)", "", ""),
            ("CARD_NUMBER", "VARCHAR2(200)", "", "Encrypted"),
            ("CARD_NUMBER_MASK", "VARCHAR2(30)", "", "Masked display"),
            ("AUTH_CODE", "VARCHAR2(200)", "", "Encrypted"),
            ("SCHEME", "VARCHAR2(20)", "", ""),
            ("MERCHANT_VPA", "VARCHAR2(100)", "", ""),
            ("GATEWAY_TXN_ID", "VARCHAR2(50)", "", ""),
            ("ARN", "VARCHAR2(50)", "", ""),
            ("IS_DELETED", "CHAR(1)", "DEFAULT N", ""),
            ("CREATED_DATE", "TIMESTAMP", "", ""),
        ],
    )

    # ----- REFUND_REQUEST -----
    doc.add_heading("REFUND_REQUEST", level=2)
    add_col_table(
        doc,
        [
            ("REQUEST_ID", "NUMBER(30)", "PK", "Request ID"),
            ("REQUEST_TYPE", "VARCHAR2(10)", "", "UPI/POS"),
            ("TXN_ID", "NUMBER(15)", "FK", "VENDOR_TRANSACTION"),
            ("CUSTOMER_ID", "VARCHAR2(20)", "", ""),
            ("ACCOUNT_NO", "VARCHAR2(20)", "", ""),
            ("BRANCH_SOL", "VARCHAR2(10)", "", ""),
            ("RRN", "VARCHAR2(30)", "", ""),
            ("MID", "VARCHAR2(30)", "", ""),
            ("TID", "VARCHAR2(30)", "", ""),
            ("TXN_DATE", "DATE", "", ""),
            ("TXN_AMOUNT", "NUMBER(15,2)", "", ""),
            ("REFUND_AMOUNT", "NUMBER(15,2)", "", ""),
            ("CARD_NUMBER_MASK", "VARCHAR2(30)", "", ""),
            ("SCHEME", "VARCHAR2(20)", "", ""),
            ("MERCHANT_VPA", "VARCHAR2(100)", "", ""),
            ("GATEWAY_TXN_ID", "VARCHAR2(50)", "", ""),
            ("VENDOR", "VARCHAR2(30)", "", ""),
            ("REMARK", "VARCHAR2(500)", "", "Latest remark"),
            ("REFERENCE_DOC_PATH", "VARCHAR2(500)", "", ""),
            ("STATUS", "VARCHAR2(40)", "", "Current Status"),
            ("STAGE", "VARCHAR2(30)", "", "BRANCH / DCO_REVIEW / DCO_APPROVE / VENDOR / FINACLE / CLOSED"),
            ("RETRY_COUNT", "NUMBER(2)", "", ""),
            ("VENDOR_REJECT_REASON", "VARCHAR2(500)", "", ""),
            ("CREATED_BY", "VARCHAR2(10)", "FK", "Branch Maker EIN — cannot approve own request"),
            ("CREATED_DATE", "TIMESTAMP", "", ""),
            ("BRANCH_APPROVED_BY", "VARCHAR2(10)", "FK", "Branch Checker EIN (<> CREATED_BY)"),
            ("BRANCH_APPROVED_DATE", "TIMESTAMP", "", ""),
            ("BRANCH_REMARK", "VARCHAR2(500)", "", ""),
            ("DCO_REVIEWED_BY", "VARCHAR2(10)", "FK", "DCO Reviewer EIN (admin or normal DCO)"),
            ("DCO_REVIEWED_DATE", "TIMESTAMP", "", ""),
            ("DCO_REVIEW_REMARK", "VARCHAR2(500)", "", ""),
            ("DCO_APPROVED_BY", "VARCHAR2(10)", "FK", "DCO Final Approver EIN (<> DCO_REVIEWED_BY)"),
            ("DCO_APPROVED_DATE", "TIMESTAMP", "", ""),
            ("DCO_APPROVE_REMARK", "VARCHAR2(500)", "", ""),
            ("UPDATED_BY", "VARCHAR2(10)", "FK", "USER"),
            ("UPDATED_DATE", "TIMESTAMP", "", ""),
            ("IS_DELETED", "CHAR(1)", "DEFAULT N", "Logical delete"),
        ],
    )

    # ----- REQUEST_HISTORY -----
    doc.add_heading("REQUEST_HISTORY", level=2)
    add_col_table(
        doc,
        [
            ("HISTORY_ID", "NUMBER(15)", "PK", ""),
            ("REQUEST_ID", "NUMBER(30)", "FK", "REFUND_REQUEST"),
            ("ACTION", "VARCHAR2(50)", "", "CREATE / BRANCH_APPROVE / DCO_REVIEW / DCO_APPROVE etc."),
            ("OLD_STATUS", "VARCHAR2(40)", "", ""),
            ("NEW_STATUS", "VARCHAR2(40)", "", ""),
            ("REMARK", "VARCHAR2(500)", "", ""),
            ("ACTION_BY", "VARCHAR2(10)", "FK", "USER EIN"),
            ("ACTION_DATE", "TIMESTAMP", "", ""),
        ],
    )

    # ----- USER_APPROVAL -----
    doc.add_heading("USER_APPROVAL", level=2)
    add_col_table(
        doc,
        [
            ("APPROVAL_ID", "NUMBER(10)", "PK", ""),
            ("REQUEST_TYPE", "VARCHAR2(20)", "", "ADD / ACTIVATE / DEACTIVATE / DELETE / MAKE_ADMIN / REVOKE_ADMIN"),
            ("TARGET_EIN", "VARCHAR2(10)", "FK", "User being acted upon"),
            ("TARGET_USERNAME", "VARCHAR2(100)", "", "Snapshot for ADD"),
            ("TARGET_EMAIL", "VARCHAR2(150)", "", "Snapshot for ADD"),
            ("TARGET_ROLE_ID", "NUMBER(5)", "FK", "Snapshot role"),
            ("TARGET_USER_TYPE", "VARCHAR2(30)", "", "Snapshot user type"),
            ("TARGET_SOL_ID", "VARCHAR2(10)", "", "Snapshot SOL"),
            ("TARGET_VALID_FROM", "DATE", "", "Snapshot — temp users"),
            ("TARGET_VALID_TILL", "DATE", "", "Snapshot — temp users"),
            ("REQUEST_ADMIN_YN", "CHAR(1)", "DEFAULT N", "Y = grant IS_ADMIN on approve"),
            ("STATUS", "VARCHAR2(20)", "", "PENDING / APPROVED / REJECTED"),
            ("MAKER_ID", "VARCHAR2(10)", "FK", "Maker EIN (must be admin for ADD / MAKE_ADMIN)"),
            ("MAKER_REMARK", "VARCHAR2(500)", "", ""),
            ("MAKER_DATE", "TIMESTAMP", "", ""),
            ("CHECKER_ID", "VARCHAR2(10)", "FK", "Checker EIN (<> MAKER_ID)"),
            ("CHECKER_REMARK", "VARCHAR2(500)", "", ""),
            ("CHECKER_DATE", "TIMESTAMP", "", "Action date"),
        ],
    )

    # ----- FILE_UPLOAD -----
    doc.add_heading("FILE_UPLOAD", level=2)
    add_col_table(
        doc,
        [
            ("FILE_ID", "NUMBER(10)", "PK", ""),
            ("FILE_NAME", "VARCHAR2(255)", "", ""),
            ("VENDOR", "VARCHAR2(30)", "", ""),
            ("FILE_TYPE", "VARCHAR2(20)", "", "POS/UPI/ARN"),
            ("SOURCE", "VARCHAR2(20)", "", "AUTO/MANUAL"),
            ("STATUS", "VARCHAR2(20)", "", ""),
            ("STORAGE_PATH", "VARCHAR2(500)", "", ""),
            ("UPLOAD_TIME", "TIMESTAMP", "", ""),
            ("UPLOADED_BY", "VARCHAR2(10)", "FK", "USER EIN"),
        ],
    )

    # ----- AUDIT_LOG -----
    doc.add_heading("AUDIT_LOG", level=2)
    add_col_table(
        doc,
        [
            ("AUDIT_ID", "NUMBER(15)", "PK", ""),
            ("APP_CODE", "VARCHAR2(10)", "", "MAT / MEA / COMMON"),
            ("USER_ID", "VARCHAR2(10)", "FK", "USER EIN"),
            ("MODULE", "VARCHAR2(50)", "", ""),
            ("ACTION", "VARCHAR2(100)", "", ""),
            ("REQUEST_ID", "NUMBER(30)", "FK", "Nullable"),
            ("IP_ADDRESS", "VARCHAR2(50)", "", ""),
            ("OLD_VALUE", "CLOB", "", ""),
            ("NEW_VALUE", "CLOB", "", ""),
            ("ACTION_TIME", "TIMESTAMP", "", ""),
        ],
    )

    # ----- SCHEDULER_LOG -----
    doc.add_heading("SCHEDULER_LOG", level=2)
    add_col_table(
        doc,
        [
            ("JOB_ID", "NUMBER(10)", "PK", ""),
            ("JOB_NAME", "VARCHAR2(100)", "", ""),
            ("START_TIME", "TIMESTAMP", "", ""),
            ("END_TIME", "TIMESTAMP", "", ""),
            ("STATUS", "VARCHAR2(20)", "", ""),
            ("MESSAGE", "VARCHAR2(1000)", "", ""),
        ],
    )

    # ----- EMAIL_LOG -----
    doc.add_heading("EMAIL_LOG (Optional)", level=2)
    add_col_table(
        doc,
        [
            ("EMAIL_ID", "NUMBER(10)", "PK", ""),
            ("REQUEST_ID", "NUMBER(30)", "FK", ""),
            ("TO_EMAIL", "VARCHAR2(200)", "", ""),
            ("SUBJECT", "VARCHAR2(300)", "", ""),
            ("STATUS", "VARCHAR2(20)", "", ""),
            ("SENT_TIME", "TIMESTAMP", "", ""),
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
